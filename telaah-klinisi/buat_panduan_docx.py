# -*- coding: utf-8 -*-
r"""Ubah panduan-penilaian.md menjadi .docx agar setara dengan buku penilaiannya.

Penilai adalah klinisi, bukan pembaca berkas markdown: tabel pipa tidak terbaca sebagai
teks biasa. Skrip ini menyalin isi panduan apa adanya ke dalam dokumen Word, memakai gaya
yang sama dengan buku penilaian sehingga keduanya terlihat satu rangkaian.

Cakupan penanda markdown sengaja dibatasi pada yang memang dipakai berkas itu: judul
(#, ##, ###), tabel, daftar bernomor dan bertitik, serta **tebal**, *miring*, dan `kode`
dalam baris. Penanda lain tidak ditafsirkan.

    python telaah-klinisi/buat_panduan_docx.py
"""
import io
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx_helper import field, heading, small, INK, MUTED  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HERE = os.path.join(BASE, "telaah-klinisi")
SRC = os.path.join(HERE, "panduan-penilaian.md")
OUT = os.path.join(HERE, "panduan-penilaian.docx")

CODE_BG = RGBColor(0x33, 0x44, 0x55)
INLINE = re.compile(r"(\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`)")


def inline(par, text, size=10.5):
    """Render **bold**, *italic* and `code` inside one paragraph."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            r = par.add_run(piece[2:-2]); r.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            r.font.name = "Consolas"; r.font.color.rgb = CODE_BG
            r.font.size = Pt(size - 1); continue
        elif piece.startswith("*") and piece.endswith("*"):
            r = par.add_run(piece[1:-1]); r.italic = True
        else:
            r = par.add_run(piece)
        r.font.size = Pt(size)


def plain(text):
    """Heading runs carry no inline formatting, so drop the markers rather than show them."""
    return text.replace("**", "").replace("`", "").replace("*", "")


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    head, body = rows[0], rows[2:]          # rows[1] is the |---|---| separator
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    t.autofit = True
    for i, h in enumerate(head):
        p = t.rows[0].cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        inline(p, h, 9.5)
        for r in p.runs:
            r.bold = True
    for row in body:
        cells = t.add_row().cells
        for i, val in enumerate(row[:len(head)]):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            inline(p, val, 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    lines = io.open(SRC, encoding="utf-8").read().split("\n")

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.8)

    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    small(foot, "Panduan Penilaian — Telaah Klinisi     ", 8)
    field(foot, "PAGE", "1")
    small(foot, " / ", 8)
    field(foot, "NUMPAGES", "1")

    i = 0
    while i < len(lines):
        line = lines[i]
        t = line.strip()

        if not t:
            i += 1
            continue

        if t.startswith("|"):                       # gather the whole table
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(split_row(lines[i]))
                i += 1
            if len(block) >= 2:
                add_table(doc, block)
            continue

        if t.startswith("### "):
            heading(doc, plain(t[4:]), level=3, size=11, space_before=12)
        elif t.startswith("## "):
            heading(doc, plain(t[3:]), level=2, size=13, space_before=16)
        elif t.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(10)
            r = p.add_run(plain(t[2:])); r.bold = True; r.font.size = Pt(19); r.font.color.rgb = INK
        elif re.match(r"^\d+\.\s", t):
            body = re.sub(r"^\d+\.\s", "", t)
            # fold any indented continuation lines into the same item
            while i + 1 < len(lines) and lines[i + 1].startswith("   ") and lines[i + 1].strip():
                i += 1
                body += " " + lines[i].strip()
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            inline(p, body)
        elif t.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            inline(p, t[2:])
        else:
            # Markdown wraps a paragraph across source lines; a blank line ends it.
            # Emitting one Word paragraph per source line breaks prose into fragments
            # with spacing between them, so gather the run first.
            body = [t]
            while i + 1 < len(lines):
                nxt = lines[i + 1].strip()
                if not nxt or nxt.startswith(("|", "#", "- ", "* ")) or re.match(r"^\d+\.\s", nxt):
                    break
                i += 1
                body.append(nxt)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            inline(p, " ".join(body))
        i += 1

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print("ditulis: %s (%.0f KB)" % (os.path.relpath(path, BASE), os.path.getsize(path) / 1024))
    print("Jalankan siapkan_kiriman.py agar ikut tersalin ke folder tiap penilai.")
