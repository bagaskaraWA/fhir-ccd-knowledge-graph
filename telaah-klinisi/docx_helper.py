# -*- coding: utf-8 -*-
"""Pembantu penyusunan dokumen Word yang dipakai bersama oleh pembangkit berkas telaah.

Sebelumnya fungsi-fungsi ini menumpang pada buat_booklet.py. Berkas itu kini sudah
digantikan rancangan penilaian per rekomendasi, sehingga pembantunya dipindahkan ke sini
agar skrip yang masih aktif tidak mengimpor dari skrip yang sudah tidak dipakai.
"""
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x60, 0x60, 0x60)


def field(par, instr, placeholder=""):
    """Sisipkan kode field Word (TOC, PAGE, NUMPAGES)."""
    run = par.add_run()
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "separate")
    if placeholder:
        ph = OxmlElement("w:t"); ph.text = placeholder; b.append(ph)
    c = OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"), "end")
    for el in (a, t, b, c):
        run._r.append(el)
    return run


def small(par, text, size=8.5, color=MUTED, italic=False):
    r = par.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.italic = italic
    return r


def table(doc, headers, rows, widths, size=8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.width = Cm(widths[i])
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h); r.bold = True; r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = Cm(widths[i])
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val)); r.font.size = Pt(size)
    return t


def heading(doc, text, level=2, size=11, space_before=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = INK
    if level:
        p.style = doc.styles["Heading %d" % level]
        for run in p.runs:
            run.font.color.rgb = INK
            run.font.size = Pt(size)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
