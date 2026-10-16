# -*- coding: utf-8 -*-
r"""Terapkan perubahan pascapembangkitan pada lembar telaah seorang penilai.

buat_telaah.py membangkitkan lembar dalam bentuk awalnya. Perubahan yang diminta sesudahnya
diterapkan di sini, DI TEMPAT, bukan dengan membangkitkan ulang: penulis menyunting lembar
itu dengan tangan (catatan data sintetis yang dipersingkat, kesimpulan pasien 1, kolom Kode
pasien 1), dan membangkitkan ulang akan menghapus suntingan tersebut.

Langkah-langkahnya, seluruhnya aman dijalankan ulang:

  1. Pemisah halaman sebelum judul pasien -> page_break_before pada judulnya, supaya teks
     yang lebih panjang tidak meninggalkan halaman kosong.
  2. Pasien yang tidak tercantum pada pasien-telaah.txt dibuang (30 -> 20).
  3. Kesimpulan generator -> terjemahan Indonesia ditambah suntingan penulis
     (kesimpulan_id.teks_akhir).
  4. Baris cek/stop -> "Cek/hentikan penggunaan <obat> pada pasien — <akibat>"
     (label_rekomendasi.teks_cek).
  5. Kolom Kode dihapus. Lebar kolom lain tidak diubah, sama seperti yang dilakukan penulis
     pada pasien 1.
  6. Pasien dinomori ulang 1..n, dan kalimat jumlah baris pada halaman pembuka diperbarui.
  7. Istilah tampilan (label_rekomendasi.ISTILAH: "beta blocker", "CKD") pada kolom
     Rekomendasi. Kesimpulan sudah memperolehnya lewat langkah 3.
  8. Kolom Penilaian -> tiga kotak centang Word (☐) yang dapat diklik, satu untuk tiap
     pilihan, dan judul kolomnya menjadi "Penilaian (centang satu)". Kotak yang sama tercetak
     sebagai ☐, jadi lembar dapat diisi di layar maupun dengan pena. Kotak yang dapat diklik
     bisa tercentang tanpa sengaja saat berkas disiapkan lalu ikut tercetak, karena itu skrip
     melaporkan jumlah kotak yang tercentang; sebelum mencetak, angka itu harus nol. Kotak
     cetak berupa huruf biasa dari versi sebelumnya dibungkus menjadi kotak centang. Hanya
     .docx.
  9. Tata letak cetak: A4 potret, satu pasien satu halaman, dengan tinggi baris minimum
     untuk tulisan tangan pada kolom saran. Hanya .docx.

Berkas .csv menerima langkah 2, 4, 6, dan 7, serta kehilangan kolom recId dan cor -- padanan
kolom Kode -- agar penilai yang mengisi lewat Excel melihat hal yang sama. Kode itu tidak
hilang: sebelum dihapus, ia disimpan ke kunci-telaah.csv, yang TIDAK dikirim ke penilai dan
menjadi satu-satunya jalan memetakan jawaban kembali ke rekomendasinya.

Teks yang tidak cocok dengan bentuk lama maupun bentuk baru dianggap suntingan tangan dan
dibiarkan, lalu dilaporkan jumlahnya. Terjemahan diperiksa secara mekanis lebih dahulu --
kutipan, target tekanan darah, ambang BMI, stadium, kelas rekomendasi, dan nama obat harus
terbawa. Pemeriksaan itu tidak menggantikan pembacaan oleh penutur.

Untuk penilai kedua, menyalin lembar R1 yang sudah final lebih aman daripada menjalankan
skrip ini pada R2: suntingan tangan pada R1 tidak tercatat di mana pun selain berkasnya.

    python telaah-klinisi/perbarui_lembar.py --periksa    # hanya pemeriksaan terjemahan
    python telaah-klinisi/perbarui_lembar.py R1 --uji     # uji kering: laporkan, jangan tulis
    python telaah-klinisi/perbarui_lembar.py R1           # terapkan pada lembar R1

Untuk MEMERIKSA lembar yang sudah disunting atau dicetak, selalu pakai --uji. Tanpa --uji,
skrip menulis ulang berkasnya. Setelah berkas disimpan Word, langkah 9 selalu melaporkan satu
perubahan walau tidak ada yang diubah: Word membuang atribut yang sama dengan bawaannya
(misalnya hRule="atLeast" pada tinggi baris dan ukuran kotak yang sama dengan gaya Normal),
dan langkah itu memasangnya kembali. Tampilannya tidak berubah, tetapi berkasnya tertulis.
"""
import argparse
import copy
import csv
import io
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import buat_telaah as B                           # noqa: E402
from kesimpulan_id import TERJEMAHAN, SUNTINGAN, teks_akhir   # noqa: E402
from label_rekomendasi import JUDUL, teks_cek, istilah   # noqa: E402

from docx import Document                         # noqa: E402
from docx.oxml import OxmlElement                 # noqa: E402
from docx.oxml.ns import qn                       # noqa: E402
from docx.text.paragraph import Paragraph         # noqa: E402

KIRIM = os.path.join(B.HERE, "kirim")
DAFTAR_TELAAH = os.path.join(B.HERE, "pasien-telaah.txt")
KUNCI = os.path.join(B.HERE, "kunci-telaah.csv")
PENGENAL = re.compile(r"pengenal ([0-9a-f]{8})\b")
JUDUL_PASIEN = re.compile(r"^\d+\.\s")
KOLOM_KODE_CSV = ("recId", "cor")

# ejaan Inggris -> ejaan Indonesia; yang tidak tercantum ditulis sama pada kedua bahasa
OBAT = {"clopidogrel": "klopidogrel", "nitroglycerin": "nitrogliserin",
        "naproxen": "naproksen", "succinate": "suksinat", "metoprolol": "metoprolol",
        "prasugrel": "prasugrel", "propranolol": "propranolol", "ibuprofen": "ibuprofen",
        "warfarin": "warfarin", "statin": "statin", "aspirin": "aspirin"}


def baca_daftar(path):
    return [l.split()[0] for l in io.open(path, encoding="utf-8")
            if l.strip() and not l.startswith("#")]


def pasien():
    return baca_daftar(B.LIST)


# ── pemeriksaan terjemahan ────────────────────────────────────────────────────────────

def unsur(teks, bahasa):
    """Unsur yang wajib terbawa, dinormalisasi agar kedua bahasa dapat dibandingkan."""
    t = teks
    u = {
        "kutipan": sorted(re.findall(r"\[[^\]]+\]", t)),
        "target": sorted(re.findall(r"<\s*\d+\s*/\s*<?\s*\d+", t)),
        "bmi": sorted(re.findall(r"BMI\s*\d+\+", t)),
        "cor": sorted(re.sub(r"\s+", "", c) for c in re.findall(r"COR\s*\d", t)),
    }
    low = t.lower()
    if bahasa == "en":
        st = re.findall(r"stages?\s+(\d)(?:\s*-\s*(\d))?", t, re.I)
        kelas = re.findall(r"Class\s+(\d)", t)
        obat = sorted(OBAT[o] for o in OBAT for _ in re.findall(r"\b%s\b" % o, low))
    else:
        st = re.findall(r"stadium\s+(\d)(?:\s*—\s*(\d))?", t)
        kelas = re.findall(r"Kelas\s+(\d)", t)
        obat = sorted(v for v in OBAT.values() for _ in re.findall(r"\b%s\b" % v, low))
    u["stadium"] = sorted(st)
    u["kelas"] = sorted(kelas)
    u["obat"] = obat
    return u


def periksa(pids):
    """Kembalikan daftar masalah terjemahan; kosong berarti seluruhnya lolos."""
    masalah = []
    for pid in pids:
        en = B.bottom_line(pid)
        idn = TERJEMAHAN.get(pid[:8])
        if not idn:
            masalah.append("%s: tidak ada terjemahan" % pid[:8])
            continue
        a, b = unsur(en, "en"), unsur(idn, "id")
        for k in a:
            if a[k] != b[k]:
                masalah.append("%s: %s berbeda\n      en=%s\n      id=%s" % (pid[:8], k, a[k], b[k]))
        # kata "never" bermuatan klinis dan mudah hilang dalam penerjemahan
        if re.search(r"\bnever\b", en, re.I) and "tidak boleh" not in idn:
            masalah.append("%s: larangan 'never' tidak terbawa" % pid[:8])
    # Suntingan penulis sengaja menyimpang dari terjemahan setia, jadi tidak diperiksa
    # terhadap sumbernya; yang diperiksa hanyalah bahwa setiap penggalannya ditemukan.
    for pid8 in SUNTINGAN:
        try:
            teks_akhir(pid8)
        except ValueError as e:
            masalah.append("suntingan: %s" % e)
    return masalah


# ── bagian-bagian dokumen ─────────────────────────────────────────────────────────────

def judul_pasien(el, doc):
    if el.tag != qn("w:p"):
        return False
    p = Paragraph(el, doc._body)
    return p.style is not None and p.style.name == "Heading 1" and bool(JUDUL_PASIEN.match(p.text))


def bagian_pasien(doc):
    """Daftar (pid8, [elemen]) untuk setiap pasien, berurutan sesuai dokumen."""
    out, cur = [], None
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:sectPr"):
            break
        if judul_pasien(el, doc):
            cur = [None, [el]]
            out.append(cur)
        elif cur is not None:
            cur[1].append(el)
            if cur[0] is None:
                teks = "".join(x.text or "" for x in el.iter(qn("w:t")))
                m = PENGENAL.search(teks)
                if m:
                    cur[0] = m.group(1)
    for pid8, _ in out:
        if pid8 is None:
            sys.exit("ada bagian pasien tanpa baris 'pengenal' -- tidak disimpan.")
    return [(p, els) for p, els in out]


def ganti_teks(par, teks):
    """Ganti isi paragraf dengan mempertahankan format run pertamanya."""
    runs = par.runs
    runs[0].text = teks
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


# ── langkah 1: pemisah halaman ────────────────────────────────────────────────────────

def rapikan_pemisah_halaman(doc):
    """Ganti paragraf pemisah halaman sebelum judul pasien dengan page_break_before.

    buat_telaah.py memisahkan pasien dengan paragraf kosong berisi page break. Bila tabel
    pasien sebelumnya tepat memenuhi halaman, paragraf itu terdorong ke halaman baru lalu
    memecah halaman sekali lagi, sehingga terbentuk halaman kosong. Terjemahan yang lebih
    panjang memicu hal itu pada pasien 23. Menaruh pemisah pada judulnya sendiri membuat
    halaman kosong tidak mungkin terjadi, tanpa mengubah teks apa pun.
    """
    n = 0
    for par in list(doc.paragraphs):
        el = par._p
        hanya_pemisah = (not par.text.strip()
                         and any(br.get(qn("w:type")) == "page" for br in el.iter(qn("w:br"))))
        if not hanya_pemisah:
            continue
        nxt = el.getnext()
        if nxt is None or not judul_pasien(nxt, doc):
            continue
        Paragraph(nxt, par._parent).paragraph_format.page_break_before = True
        el.getparent().remove(el)
        n += 1
    return n


# ── langkah 2: buang pasien ───────────────────────────────────────────────────────────

def pangkas_pasien(doc, simpan8):
    n = 0
    for pid8, els in bagian_pasien(doc):
        if pid8 in simpan8:
            continue
        for el in els:
            el.getparent().remove(el)
        n += 1
    hadir = [p for p, _ in bagian_pasien(doc)]
    if sorted(hadir) != sorted(simpan8):
        sys.exit("sesudah pemangkasan, pasien yang ada (%d) tidak sama dengan daftar telaah "
                 "(%d) -- tidak disimpan." % (len(hadir), len(simpan8)))
    return n


# ── langkah 3: terjemahan kesimpulan ──────────────────────────────────────────────────

def terjemahkan(doc, n_pasien):
    """Kembalikan (diganti, sudah, tangan).

    Sasaran tiap kesimpulan adalah teks_akhir(): terjemahan setia, suntingan penulis, lalu
    istilah tampilan. Bentuk yang dikenali lalu diganti adalah teks Inggris asli, terjemahan
    setia, dan bentuk bersunting tetapi belum beristilah tampilan; teks lain dianggap
    suntingan tangan dan dibiarkan.
    """
    akhir = {teks_akhir(p) for p in pasien()}
    ke_akhir = {}
    for p in pasien():
        for bentuk in (B.bottom_line(p), TERJEMAHAN[p[:8]], teks_akhir(p, dengan_istilah=False)):
            ke_akhir[bentuk] = teks_akhir(p)
    ps = doc.paragraphs
    diganti = sudah = tangan = 0
    for i, par in enumerate(ps):
        if par.text != "Kesimpulan generator" or i + 1 >= len(ps):
            continue
        q = ps[i + 1]
        teks = "".join(r.text for r in q.runs)
        if teks in akhir:
            sudah += 1
        elif teks in ke_akhir:
            ganti_teks(q, ke_akhir[teks])
            diganti += 1
        else:
            tangan += 1
    if diganti + sudah + tangan != n_pasien:
        sys.exit("kesimpulan: %d ditemukan untuk %d pasien -- tidak disimpan."
                 % (diganti + sudah + tangan, n_pasien))
    return diganti, sudah, tangan


# ── langkah 4: baris cek/stop ─────────────────────────────────────────────────────────

def peta_cek():
    """Teks lama -> teks baru (beristilah tampilan) untuk setiap baris cek/stop."""
    import json
    peta = {}
    for pid in pasien():
        fp = os.path.join(B.EVAL, "briefing_facts_%s.json" % pid)
        facts = json.load(io.open(fp, encoding="utf-8"))
        for kel, rid, cor, teks, obat in B.rows_for(facts):
            if kel == "cek":
                lama = ("%s — %s" % (obat, teks)) if obat else teks   # sama dengan buat_telaah
                baru = istilah(teks_cek(rid, obat))
                peta[lama] = baru
                if teks_cek(rid, obat) != baru:
                    peta[teks_cek(rid, obat)] = baru
    return peta


def cek_docx(doc, peta):
    baru_set = set(peta.values())
    diganti = sudah = tangan = 0
    for t in doc.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if "Rekomendasi" not in hdr:
            continue
        k = hdr.index("Rekomendasi")
        for row in t.rows[1:]:
            cells = row.cells
            if cells[0].text.strip() != B.SINGKAT["cek"]:
                continue
            par = cells[k].paragraphs[0]
            teks = "".join(r.text for r in par.runs)
            if teks in baru_set:
                sudah += 1
            elif teks in peta:
                ganti_teks(par, peta[teks])
                diganti += 1
            else:
                tangan += 1
    return diganti, sudah, tangan


# ── langkah 5: kolom Kode ─────────────────────────────────────────────────────────────

def hapus_kolom_kode(doc):
    n = 0
    for t in doc.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if "Kode" not in hdr:
            continue                           # misalnya pasien 1, yang sudah disunting tangan
        k = hdr.index("Kode")
        for tr in t._tbl.findall(qn("w:tr")):
            tr.remove(tr.findall(qn("w:tc"))[k])
        grid = t._tbl.tblGrid
        grid.remove(grid.findall(qn("w:gridCol"))[k])
        n += 1
    return n


# ── langkah 6: nomor ulang dan jumlah baris ───────────────────────────────────────────

def nomori_ulang(doc):
    n = diubah = 0
    for par in doc.paragraphs:
        if not (par.style is not None and par.style.name == "Heading 1"
                and JUDUL_PASIEN.match(par.text)):
            continue
        n += 1
        r0 = par.runs[0]
        if not re.match(r"^\d+", r0.text):
            sys.exit("nomor judul pasien tidak berada pada run pertamanya -- tidak disimpan.")
        baru = re.sub(r"^\d+", str(n), r0.text, count=1)
        if baru != r0.text:
            r0.text = baru
            diubah += 1
    return diubah


def istilah_sel(doc):
    """Langkah 7: istilah tampilan pada setiap sel kolom Rekomendasi."""
    n = 0
    for t in doc.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if "Rekomendasi" not in hdr:
            continue
        k = hdr.index("Rekomendasi")
        for row in t.rows[1:]:
            par = row.cells[k].paragraphs[0]
            teks = "".join(r.text for r in par.runs)
            if istilah(teks) != teks:
                ganti_teks(par, istilah(teks))
                n += 1
    return n


# ── langkah 8: kotak centang ──────────────────────────────────────────────────────────

PILIHAN = ("Tepat", "Belum tepat", "Tidak dapat dinilai")
JUDUL_PENILAIAN = "Penilaian (centang satu)"
KOTAK = "☐"                       # ☐, dari MS Gothic seperti kotak centang Word
TERCENTANG = "☒"                  # ☒, tampilan kotak centang Word yang dicentang


def run_kotak():
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:eastAsia", "w:hAnsi"):
        f.set(qn(a), "MS Gothic")
    f.set(qn("w:hint"), "eastAsia")
    rpr.append(f)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = KOTAK
    r.append(t)
    return r


def isi_kotak(par):
    """Ganti teks pilihan dengan kotak yang masing-masing melekat pada labelnya.

    Spasi tak-putus di antara kotak dan labelnya -- dan di dalam label -- mencegah kotak
    terpisah dari labelnya saat baris membungkus; pembungkusan hanya terjadi di antara
    pilihan, sehingga sel tetap dua baris seperti semula.
    """
    runs = par.runs
    templat = copy.deepcopy(runs[0]._r.rPr) if runs[0]._r.rPr is not None else None
    for r in runs:
        r._r.getparent().remove(r._r)
    for k, label in enumerate(PILIHAN):
        par._p.append(run_kotak())
        run = OxmlElement("w:r")
        if templat is not None:
            run.append(copy.deepcopy(templat))
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = " " + label.replace(" ", " ") + ("  " if k < len(PILIHAN) - 1 else "")
        run.append(t)
        par._p.append(run)


def kotak_centang_par(par):
    """Kotak centang Word (w14:checkbox) di dalam paragraf."""
    return [s for s in par._p.findall(".//" + qn("w:sdt"))
            if s.find(qn("w:sdtPr") + "/" + qn("w14:checkbox")) is not None]


def bungkus_kotak(r, sdt_id):
    """Bungkus run berisi ☐ menjadi kotak centang Word; format run-nya dipertahankan."""
    sdt, pr, isi = OxmlElement("w:sdt"), OxmlElement("w:sdtPr"), OxmlElement("w:sdtContent")
    if r.find(qn("w:rPr")) is not None:
        pr.append(copy.deepcopy(r.find(qn("w:rPr"))))
    i = OxmlElement("w:id")
    i.set(qn("w:val"), str(sdt_id))
    pr.append(i)
    cb = OxmlElement("w14:checkbox")
    for tag, nilai in (("w14:checked", "0"), ("w14:checkedState", "2612"),
                       ("w14:uncheckedState", "2610")):
        el = OxmlElement(tag)
        el.set(qn("w14:val"), nilai)
        if tag != "w14:checked":
            el.set(qn("w14:font"), "MS Gothic")
        cb.append(el)
    pr.append(cb)
    sdt.append(pr)
    sdt.append(isi)
    r.addprevious(sdt)
    isi.append(r)


def kotak_penilaian(doc):
    """Kembalikan (diubah, sudah, tangan, judul_diubah, kotak_dibungkus, tercentang)."""
    diubah = sudah = tangan = judul = dibungkus = tercentang = 0
    # pengenal kontrol harus unik; mulai jauh di atas pengenal lain yang sudah ada
    sdt_id = max([int(x.get(qn("w:val"))) for x in doc.element.body.iter(qn("w:id"))
                  if x.getparent().tag == qn("w:sdtPr")] + [600000000])
    for t in doc.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if "Rekomendasi" not in hdr:
            continue
        k = next(i for i, h in enumerate(hdr) if h.startswith("Penilaian"))
        hp = t.rows[0].cells[k].paragraphs[0]
        if hp.text.strip() == "Penilaian":
            hp.runs[0].text = JUDUL_PENILAIAN
            for r in hp.runs[1:]:
                r._r.getparent().remove(r._r)
            judul += 1
        for row in t.rows[1:]:
            par = row.cells[k].paragraphs[0]
            teks = "".join(x.text or "" for x in par._p.iter(qn("w:t")))
            ada = kotak_centang_par(par)
            tercentang += sum(1 for s in ada
                              if s.find(".//" + qn("w14:checked")).get(qn("w14:val")) in ("1", "true"))
            if len(ada) == len(PILIHAN):
                sudah += 1
                continue
            if teks == B.PILIHAN:
                isi_kotak(par)
            elif ada or teks.count(KOTAK) != len(PILIHAN):
                tangan += 1
                continue
            for r in par._p.findall(qn("w:r")):           # run langsung, belum terbungkus
                if "".join(x.text or "" for x in r.iter(qn("w:t"))) == KOTAK:
                    sdt_id += 1
                    bungkus_kotak(r, sdt_id)
                    dibungkus += 1
            diubah += 1
    return diubah, sudah, tangan, judul, dibungkus, tercentang


# ── langkah 9: tata letak cetak ───────────────────────────────────────────────────────

HALAMAN_CM = dict(lebar=21.0, tinggi=29.7, kiri=1.5, kanan=1.5, atas=1.5, bawah=1.3)
LEBAR_KOLOM_CM = (2.4, 6.9, 3.7, 5.0)     # Kelompok, Rekomendasi, Penilaian, saran
PT_TABEL = 8
PT_KOTAK = 10
TINGGI_BARIS_MIN_CM = 0.9                 # ruang tulisan tangan pada kolom saran


def tata_letak(doc, halaman=None, kolom=None, pt_tabel=None, pt_kotak=None, tinggi=None):
    """Terapkan tata letak cetak. Mengembalikan 1 bila ada yang berubah, 0 bila tidak."""
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.shared import Cm, Pt
    from docx.text.run import Run
    from lxml import etree
    halaman = halaman or HALAMAN_CM
    kolom = kolom or LEBAR_KOLOM_CM
    pt_tabel = pt_tabel or PT_TABEL
    pt_kotak = pt_kotak or PT_KOTAK
    tinggi = tinggi or TINGGI_BARIS_MIN_CM

    tabel = [t for t in doc.tables
             if "Rekomendasi" in [c.text.strip() for c in t.rows[0].cells]]
    sebelum = etree.tostring(doc.sections[0]._sectPr) + b"".join(etree.tostring(t._tbl) for t in tabel)

    s = doc.sections[0]
    s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = Cm(halaman["lebar"]), Cm(halaman["tinggi"])
    s.left_margin, s.right_margin = Cm(halaman["kiri"]), Cm(halaman["kanan"])
    s.top_margin, s.bottom_margin = Cm(halaman["atas"]), Cm(halaman["bawah"])

    for t in tabel:
        if len(t._tbl.tblGrid.findall(qn("w:gridCol"))) != len(kolom):
            sys.exit("tabel tidak berkolom %d -- tidak disimpan." % len(kolom))
        for g, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), kolom):
            g.set(qn("w:w"), str(round(w * 567)))
        for i, row in enumerate(t.rows):
            for tc, w in zip(row._tr.findall(qn("w:tc")), kolom):
                tw = tc.get_or_add_tcPr().get_or_add_tcW()
                tw.set(qn("w:w"), str(round(w * 567)))
                tw.set(qn("w:type"), "dxa")
            trpr = row._tr.get_or_add_trPr()
            if trpr.find(qn("w:cantSplit")) is None:
                trpr.insert(0, OxmlElement("w:cantSplit"))
            if i > 0:
                row.height, row.height_rule = Cm(tinggi), WD_ROW_HEIGHT_RULE.AT_LEAST
        for r in t._tbl.iter(qn("w:r")):
            teks = "".join(x.text or "" for x in r.iter(qn("w:t")))
            Run(r, None).font.size = Pt(pt_kotak if teks in (KOTAK, TERCENTANG) else pt_tabel)

    sesudah = etree.tostring(doc.sections[0]._sectPr) + b"".join(etree.tostring(t._tbl) for t in tabel)
    return int(sebelum != sesudah)


def perbarui_jumlah(doc):
    n_pasien = len(bagian_pasien(doc))
    n_baris = sum(len(t.rows) - 1 for t in doc.tables
                  if "Rekomendasi" in [c.text.strip() for c in t.rows[0].cells])
    kalimat = "Jumlah baris penilaian: %d, tersebar pada %d pasien." % (n_baris, n_pasien)
    for par in doc.paragraphs:
        if par.text.startswith("Jumlah baris penilaian:"):
            if par.text != kalimat:
                ganti_teks(par, kalimat)
                return 1, n_baris, n_pasien
            return 0, n_baris, n_pasien
    sys.exit("kalimat jumlah baris tidak ditemukan -- tidak disimpan.")


# ── berkas .csv ───────────────────────────────────────────────────────────────────────

def perbarui_csv(path, urutan8, peta):
    """urutan8: pid8 menurut urutan akhir pada .docx. Kembalikan ringkasan perubahan."""
    raw = open(path, "rb").read()
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"), newline="")))
    hdr, body = rows[0], rows[1:]
    iP, iNo, iK, iR = (hdr.index(c) for c in ("pasien", "no", "kelompok", "rekomendasi"))
    nomor = {p: i + 1 for i, p in enumerate(urutan8)}
    h = {}

    sisa = [r for r in body if r[iP][:8] in nomor]
    h["baris dibuang"] = len(body) - len(sisa)
    h["nomor diubah"] = 0
    for r in sisa:
        baru = str(nomor[r[iP][:8]])
        if r[iNo] != baru:
            r[iNo] = baru
            h["nomor diubah"] += 1

    baru_set = set(peta.values())
    h["cek/stop diubah"] = h["cek/stop tangan"] = 0
    for r in sisa:
        if r[iK] != JUDUL["cek"] or r[iR] in baru_set:
            continue
        if r[iR] in peta:
            r[iR] = peta[r[iR]]
            h["cek/stop diubah"] += 1
        else:
            h["cek/stop tangan"] += 1

    h["istilah diubah"] = 0
    for r in sisa:
        if istilah(r[iR]) != r[iR]:
            r[iR] = istilah(r[iR])
            h["istilah diubah"] += 1

    # Kunci ikut diselaraskan bila sudah ada, agar teks rekomendasinya tetap sama persis
    # dengan yang dilihat penilai.
    h["kunci diselaraskan"] = 0
    if os.path.exists(KUNCI):
        kr = list(csv.reader(io.open(KUNCI, encoding="utf-8-sig", newline="")))
        jR = kr[0].index("rekomendasi")
        for r in kr[1:]:
            if istilah(r[jR]) != r[jR]:
                r[jR] = istilah(r[jR])
                h["kunci diselaraskan"] += 1
        if h["kunci diselaraskan"]:
            tulis_csv(KUNCI, kr)

    # Kolom kode disimpan ke kunci lebih dahulu, baru dihapus dari berkas penilai.
    buang = [hdr.index(c) for c in KOLOM_KODE_CSV if c in hdr]
    h["kolom kode dihapus"] = len(buang)
    if buang:
        kunci_hdr = ["no", "nama", "pasien", "kelompok", "recId", "cor", "rekomendasi"]
        ik = [hdr.index(c) for c in kunci_hdr]
        tulis_csv(KUNCI, [kunci_hdr] + [[r[i] for i in ik] for r in sisa])
        hdr = [c for i, c in enumerate(hdr) if i not in buang]
        sisa = [[v for i, v in enumerate(r) if i not in buang] for r in sisa]

    if any(h[k] for k in ("baris dibuang", "nomor diubah", "cek/stop diubah",
                          "istilah diubah", "kolom kode dihapus")):
        tulis_csv(path, [hdr] + sisa)
    h["baris akhir"] = len(sisa)
    return h


UJI = False       # --uji: seluruh langkah dijalankan di memori, tidak ada berkas yang ditulis


def tulis_csv(path, rows):
    if UJI:
        return
    buf = io.StringIO(newline="")
    csv.writer(buf).writerows(rows)                    # dialek yang sama dengan buat_telaah
    open(path, "wb").write(("﻿" + buf.getvalue()).encode("utf-8"))


# ── utama ─────────────────────────────────────────────────────────────────────────────

def perbarui(rater):
    docx_path = os.path.join(KIRIM, rater, "telaah-%s.docx" % rater)
    csv_path = os.path.join(KIRIM, rater, "telaah-%s.csv" % rater)
    simpan8 = [p[:8] for p in baca_daftar(DAFTAR_TELAAH)]
    peta = peta_cek()

    # Seluruh langkah .docx dijalankan di memori lebih dahulu; sys.exit di dalamnya
    # membatalkan tanpa menulis berkas apa pun.
    doc = Document(docx_path)
    h = {}
    h["pemisah halaman"] = rapikan_pemisah_halaman(doc)
    h["pasien dibuang"] = pangkas_pasien(doc, simpan8)
    urutan8 = [p for p, _ in bagian_pasien(doc)]
    h["kesimpulan"] = terjemahkan(doc, len(urutan8))
    h["cek/stop docx"] = cek_docx(doc, peta)
    h["kolom Kode dihapus"] = hapus_kolom_kode(doc)
    h["nomor judul diubah"] = nomori_ulang(doc)
    ubah_jumlah, n_baris, n_pasien = perbarui_jumlah(doc)
    h["kalimat jumlah"] = ubah_jumlah
    h["istilah (sel)"] = istilah_sel(doc)
    kp = kotak_penilaian(doc)
    h["kotak centang"] = kp[:3]
    h["judul Penilaian"] = kp[3]
    h["kotak dibungkus"] = kp[4]
    h["kotak tercentang"] = kp[5]
    h["tata letak cetak"] = tata_letak(doc)

    ada_perubahan = (h["pemisah halaman"] or h["pasien dibuang"] or h["kesimpulan"][0]
                     or h["cek/stop docx"][0] or h["kolom Kode dihapus"]
                     or h["nomor judul diubah"] or ubah_jumlah or h["istilah (sel)"]
                     or kp[0] or kp[3] or kp[4] or h["tata letak cetak"])
    hc = perbarui_csv(csv_path, urutan8, peta)
    if hc["baris akhir"] != n_baris:
        sys.exit("csv memuat %d baris, docx %d -- .docx tidak disimpan." % (hc["baris akhir"], n_baris))
    if ada_perubahan:
        if not UJI:
            doc.save(docx_path)
    return h, hc, n_baris, n_pasien


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("rater", nargs="?", help="kode penilai, misalnya R1")
    ap.add_argument("--periksa", action="store_true", help="hanya pemeriksaan, tanpa menulis")
    ap.add_argument("--uji", action="store_true",
                    help="jalankan seluruh langkah pada lembar tanpa menulis berkas apa pun")
    a = ap.parse_args()

    masalah = periksa(pasien())
    print("pemeriksaan terjemahan: %d kesimpulan, %d masalah" % (len(pasien()), len(masalah)))
    for m in masalah:
        print("  - " + m)
    if masalah:
        return 1
    if a.periksa or not a.rater:
        return 0

    global UJI
    UJI = a.uji
    if UJI:
        print("UJI KERING -- tidak ada berkas yang ditulis; angka di bawah adalah yang AKAN diubah.")
    h, hc, n_baris, n_pasien = perbarui(a.rater)
    print("%s.docx -> %d pasien, %d baris penilaian" % (a.rater, n_pasien, n_baris))
    for k, v in h.items():
        if isinstance(v, tuple):
            v = "%d diubah, %d sudah, %d suntingan tangan dibiarkan" % v
        print("  %-20s %s" % (k, v))
    print("%s.csv" % a.rater)
    for k, v in hc.items():
        print("  %-20s %s" % (k, v))
    return 0


if __name__ == "__main__":
    sys.exit(main())
