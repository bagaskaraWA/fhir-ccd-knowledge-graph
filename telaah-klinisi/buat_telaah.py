# -*- coding: utf-8 -*-
r"""Rakit lembar telaah klinisi: satu baris untuk satu rekomendasi.

Menggantikan rancangan penilaian mutu ringkasan. Yang dinilai bukan lagi bagaimana model
bahasa menyajikan informasi, melainkan apakah rekomendasi yang dipilih sistem TEPAT bagi
pasien tersebut. Karena itu isi tabel diambil dari muatan fakta hasil penelusuran
(briefing_facts_<pid>.json), bukan dari prosa ringkasan -- sehingga artefak yang dinilai
bersifat deterministik dan dapat direproduksi.

Tiap pasien disajikan sebagai:

    1. Kesimpulan generator -- bagian BOTTOM LINE ringkasan; ATAS KEPUTUSAN PEMBIMBING
                          KLINIS, hanya inilah dasar penilaian. Vinyet klinis tidak
                          ditampilkan, sehingga penilai menilai apa yang benar-benar
                          disampaikan sistem, bukan rekam medis di baliknya.
    2. Tabel penilaian -- satu baris per rekomendasi, tiga pilihan penilaian.

Cakupan baris: seluruh safetyFlags, ditambah applicable dengan extraCount > 0 (yaitu yang
dipicu kondisi pasien, bukan perawatan CCD baku yang sama bagi semua orang).

    python telaah-klinisi/buat_telaah.py
"""
import csv
import io
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import buat_rekam_medis as B  # noqa: E402
from label_rekomendasi import JUDUL, OBS_KUNCI, URUT, label  # noqa: E402
from docx_helper import field, heading, small, table, page_break, INK, MUTED  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HERE = os.path.join(BASE, "telaah-klinisi")
EVAL = os.path.join(BASE, "eval")
LIST = os.path.join(HERE, "pasien.txt")
OUT = os.path.join(HERE, "kirim")
PENILAI = ("R1", "R2")

PILIHAN = "Tepat  /  Belum tepat  /  Tidak dapat dinilai"
SINGKAT = {"cek": "Cek/stop", "tata": "Tatalaksana", "obat": "Medikamentosa"}
KOLOM = ["penilai", "no", "nama", "pasien", "kelompok", "recId", "cor", "rekomendasi",
         "penilaian", "saran_bila_belum_tepat"]


# Judul BOTTOM LINE ditulis model bahasa dengan bentuk yang berbeda-beda antarproses --
# bernomor atau tidak, berti­tik dua atau tidak, dan isinya kadang menyambung pada baris
# yang sama dengan judulnya. Keempat ragam berikut semuanya muncul pada 30 ringkasan:
#
#     **1. BOTTOM LINE**            (isi pada baris berikutnya)
#     **BOTTOM LINE**               (tanpa nomor)
#     **BOTTOM LINE:** Pasien ...   (isi menyambung pada baris yang sama)
#     **Bottom line:** Pasien ...   (huruf kecil)
#
# Pola lama hanya mengenali ragam pertama, sehingga 8 dari 30 pasien kehilangan bagian ini.
BOTTOM = re.compile(
    r"\*\*\s*(?:\d+\.\s*)?BOTTOM\s+LINE\s*:?\s*\*\*\s*:?\s*(.*?)(?=\n\s*(?:\*\*|#)|\Z)",
    re.S | re.I)


def bottom_line(pid):
    """Ambil bagian BOTTOM LINE dari ringkasan; itulah konteks yang diminta pembimbing."""
    p = os.path.join(EVAL, "briefing_%s.md" % pid)
    if not os.path.exists(p):
        return ""
    m = BOTTOM.search(io.open(p, encoding="utf-8").read())
    if not m:
        return ""
    teks = re.sub(r"\s+", " ", m.group(1)).strip()
    # Model kadang menebalkan frasa di tengah kalimat. Paragraf Word menerima teks polos,
    # sehingga penanda bintangnya akan tercetak apa adanya bila tidak dibuang.
    return re.sub(r"\*{1,2}([^*]+)\*{1,2}", r"\1", teks)


def vinyet(pid, facts):
    """Vinyet klinis dari data terstruktur: demografi, diagnosis, obat, observasi syarat.

    Diagnosis diurutkan (disorder) dan (situation) lebih dahulu, lalu (finding), supaya
    yang paling bermakna secara klinis terbaca duluan. Tidak ada yang dibuang.
    """
    s = B.sections(json.load(io.open(B.bundle_for(pid), encoding="utf-8")))
    d = dict(s["demografi"])

    conds = [c["name"] for c in facts.get("conditions", [])]
    rank = lambda c: (0 if "(disorder)" in c or "(situation)" in c else 1, c.lower())
    conds = sorted(conds, key=rank)

    meds = [m["name"] for m in facts.get("medications", [])]

    # Setiap observasi syarat selalu muncul; yang tidak pernah diukur dinyatakan apa
    # adanya, sebab ketiadaannya sendiri kerap menjadi alasan sebuah rekomendasi menyala.
    obs = []
    for tampil, kunci in OBS_KUNCI:
        cocok = [r for r in s["observasi"] if any(k in r[0].lower() for k in kunci)]
        if not cocok:
            obs.append("%s tidak tersedia" % tampil)
            continue
        if tampil == "Tekanan darah":
            sis = next((r for r in cocok if "systolic" in r[0].lower()), None)
            dia = next((r for r in cocok if "diastolic" in r[0].lower()), None)
            if sis and dia:
                obs.append("Tekanan darah %s/%s mmHg" % (sis[7], dia[7]))
                continue
        r = cocok[0]
        obs.append("%s %s%s" % (tampil, r[7], "" if r[1] == "-" else " " + r[1]))
    return {"demografi": d, "kondisi": conds, "obat": meds, "observasi": obs,
            "nama": s["nama"]}


def rows_for(facts):
    """Baris penilaian: safetyFlags + applicable ber-extraCount>0, terurut kelompok."""
    out = []
    for sf in facts.get("safetyFlags", []):
        lab = label(sf["recId"])
        if lab:
            out.append((lab[0], sf["recId"].replace("CCD-2023:", ""),
                        str(sf.get("cor")), lab[1], sf.get("drug", "")))
    seen = set()
    for a in facts.get("applicable", []):
        if a.get("extraCount", 0) <= 0:
            continue
        lab = label(a["recId"])
        if lab and a["recId"] not in seen:
            seen.add(a["recId"])
            out.append((lab[0], a["recId"].replace("CCD-2023:", ""),
                        str(a.get("cor")), lab[1], ""))
    out.sort(key=lambda r: (URUT[r[0]], r[1]))
    return out


def build_docx(rater, data):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)      # A4 landscape
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = Cm(1.5); sec.bottom_margin = Cm(1.3)

    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    small(foot, "Telaah Klinisi — Penilai %s     " % rater, 8)
    field(foot, "PAGE", "1"); small(foot, " / ", 8); field(foot, "NUMPAGES", "1")

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(40)
    r = p.add_run("Telaah Ketepatan Rekomendasi"); r.bold = True; r.font.size = Pt(20)
    p = doc.add_paragraph()
    r = p.add_run("Sistem Rekomendasi Klinis Penyakit Jantung Koroner Kronis")
    r.font.size = Pt(12); r.font.color.rgb = MUTED
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Penilai: %s" % rater); r.bold = True; r.font.size = Pt(13)

    total = sum(len(d["baris"]) for d in data)
    for line in (
        "Untuk tiap pasien tersedia kesimpulan yang disusun sistem, diikuti daftar "
        "rekomendasi yang dikeluarkan sistem bagi pasien tersebut.",
        "",
        "Tugas Anda: untuk tiap baris rekomendasi, tentukan apakah rekomendasi itu "
        "TEPAT bagi pasien ini. Bila belum tepat, tuliskan apa yang akan Anda sarankan.",
        "",
        "Pilih “Tidak dapat dinilai” bila informasi yang tersedia memang belum "
        "cukup untuk memutuskan. Pilihan itu merupakan jawaban yang sah, bukan kegagalan.",
        "",
        "Penilaian didasarkan pada kesimpulan tersebut. Rekam medis lengkapnya tidak "
        "disertakan, sehingga bila kesimpulan itu belum memuat keterangan yang Anda "
        "perlukan, pilihan ketiga memang jawaban yang tepat.",
        "",
        "Rekam medis bersifat SINTETIS (Synthea), bukan pasien nyata, dan tidak memuat "
        "keluhan subjektif yang bermakna: keluhan utamanya statis sepanjang hidup pasien "
        "dan tidak memuat gambaran angina.",
        "",
        "Mohon TIDAK berdiskusi dengan penilai lain sebelum seluruh penilaian selesai.",
        "",
        "Jumlah baris penilaian: %d, tersebar pada %d pasien." % (total, len(data)),
    ):
        q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(4)
        if line:
            q.add_run(line).font.size = Pt(10.5)

    for i, d in enumerate(data, 1):
        page_break(doc)
        v = d["vinyet"]
        heading(doc, "%d. %s" % (i, v["nama"]), level=1, size=15, space_before=0)
        dem = v["demografi"]
        small(doc.add_paragraph(), "%s · %s · pengenal %s"
              % (dem.get("Jenis kelamin", "-"), dem.get("Usia", "-"), d["pid"][:8]), 9)

        # Vinyet klinis sengaja TIDAK ditampilkan: atas keputusan pembimbing klinis,
        # penilaian dilakukan semata-mata atas kesimpulan yang disusun sistem. vinyet()
        # tetap dihitung karena menjadi sumber nama dan demografi, dan datanya diperlukan
        # pada tahap analisis untuk menafsirkan jawaban "Tidak dapat dinilai" -- misalnya
        # membedakan penilai yang ragu dari LVEF yang memang tidak pernah diukur.
        if d["bottom"]:
            heading(doc, "Kesimpulan generator", level=2, size=11)
            q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(2)
            q.add_run(d["bottom"]).font.size = Pt(9)

        heading(doc, "Penilaian rekomendasi (%d)" % len(d["baris"]), level=2, size=11)
        rows = []
        for kel, rid, cor, teks, obat in d["baris"]:
            nama_rek = ("%s — %s" % (obat, teks)) if obat else teks
            rows.append((SINGKAT[kel], "%s\nCOR %s" % (rid, cor), nama_rek, PILIHAN, ""))
        table(doc, ["Kelompok", "Kode", "Rekomendasi", "Penilaian",
                    "Bila belum tepat / tidak dapat dinilai: saran Anda"],
              rows, [3.6, 1.9, 8.8, 5.0, 7.2], size=8.5)

    path = os.path.join(OUT, rater, "telaah-%s.docx" % rater)
    if not os.path.isdir(os.path.dirname(path)):
        os.makedirs(os.path.dirname(path))
    doc.save(path)
    return path


def main():
    pids = [l.split()[0] for l in io.open(LIST, encoding="utf-8")
            if l.strip() and not l.startswith("#")]

    data = []
    for pid in pids:
        fp = os.path.join(EVAL, "briefing_facts_%s.json" % pid)
        if not os.path.exists(fp):
            sys.exit("muatan fakta belum ada: %s" % fp)
        facts = json.load(io.open(fp, encoding="utf-8"))
        data.append({"pid": pid, "facts": facts, "vinyet": vinyet(pid, facts),
                     "bottom": bottom_line(pid), "baris": rows_for(facts)})

    total = sum(len(d["baris"]) for d in data)
    for rater in PENILAI:
        path = build_docx(rater, data)
        csv_path = os.path.join(OUT, rater, "telaah-%s.csv" % rater)
        with io.open(csv_path, "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh); w.writerow(KOLOM)
            for i, d in enumerate(data, 1):
                for kel, rid, cor, teks, obat in d["baris"]:
                    w.writerow([rater, i, d["vinyet"]["nama"], d["pid"], JUDUL[kel],
                                rid, cor, ("%s — %s" % (obat, teks)) if obat else teks,
                                "", ""])
        print("%s : %s  (%.1f MB) + %s"
              % (rater, os.path.relpath(path, BASE), os.path.getsize(path) / 1e6,
                 os.path.basename(csv_path)))

    per = sorted(len(d["baris"]) for d in data)
    print("\npasien %d | baris penilaian %d | per pasien min %d, median %d, maks %d"
          % (len(data), total, per[0], per[len(per) // 2], per[-1]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
